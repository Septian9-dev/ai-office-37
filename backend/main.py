import os
import base64
import io
import requests
import edge_tts
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, BackgroundTasks
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional
from groq import Groq
from google import genai
from google.genai import types
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from supabase import create_client, Client

from services.scheduler import start_automated_schedule, run_short_production_pipeline

load_dotenv()

app = FastAPI(title="37 AI Agents HQ Backend")

# Inisialisasi Klien Supabase untuk Logging Realtime
SUPABASE_URL = os.environ.get("SUPABASE_URL", "")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY", "")
supabase_client: Optional[Client] = (
    create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None
)

def log_agent_activity(agent_id: str, agent_name: str, division: str, text: str):
    """Menyimpan log aktivitas riil agen ke Supabase untuk disimak oleh Frontend Realtime"""
    if supabase_client:
        try:
            supabase_client.table("agent_logs").insert({
                "agent_id": agent_id,
                "agent_name": agent_name,
                "division": division,
                "text": text[:180]  # Ringkasan log untuk UI activity stream
            }).execute()
        except Exception as e:
            print(f"⚠️ Gagal menyimpan log aktivitas ke Supabase: {e}")

@app.on_event("startup")
def startup_event():
    start_automated_schedule()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Inisialisasi Klien Groq dan Gemini API secara bersamaan
groq_client = Groq(api_key=os.environ.get("GROQ_API_KEY")) if os.environ.get("GROQ_API_KEY") else None
gemini_client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")) if (os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")) else None

class ChatRequest(BaseModel):
    agent_id: str
    agent_name: str
    role: str
    division: str
    message: str
    gender: str = "male"
    image_base64: Optional[str] = None
    auto_dispatch: Optional[bool] = False
    provider: Optional[str] = "groq"

class DocRequest(BaseModel):
    content: str
    doc_type: str = "docx"

DIVISION_PROMPTS = {
    "Executive": "Berbicara tegas, visioner, berbobot kepemimpinan, dan fokus pada keputusan strategis serta keputusan bisnis.",
    "Technology": "Bicara lugas, teknis, santai, solutif, serta menguasai arsitektur sistem dan pemrograman.",
    "Data & AI": "Berpikir analitis, terukur, berbasis data kuantitatif, dan menyampaikan argumen secara terstruktur.",
    "Marketing": "Bicara antusias, kreatif, ramah, persuasif, dan fokus pada pertumbuhan brand serta eksekusi campaign.",
    "Operations": "Bicara empati, efisien, teratur, profesional, dan fokus pada kelancaran alur kerja operasional.",
    "Product": "Fokus pada pengalaman pengguna, prioritas fitur, dan metodologi pengembangan produk."
}

def send_to_telegram(text: str):
    token = os.environ.get("TELEGRAM_BOT_TOKEN")
    chat_id = os.environ.get("TELEGRAM_CHAT_ID")
    if not token or not chat_id:
        return False
    url = f"https://api.telegram.org/bot{token}/sendMessage"
    payload = {"chat_id": chat_id, "text": f"🚀 [HQ Mastermind Report]\n\n{text}", "parse_mode": "Markdown"}
    try:
        res = requests.post(url, json=payload)
        return res.status_code == 200
    except Exception as e:
        print(f"Telegram Error: {e}")
        return False

def send_to_whatsapp(text: str):
    wa_token = os.environ.get("WHATSAPP_API_TOKEN")
    target = os.environ.get("WHATSAPP_TARGET_PHONE")
    if not wa_token or not target:
        return False
    headers = {"Authorization": wa_token}
    payload = {"target": target, "message": f"🚀 *[HQ Mastermind Report]*\n\n{text}"}
    try:
        res = requests.post("https://api.fonnte.com/send", headers=headers, data=payload)
        return res.status_code == 200
    except Exception as e:
        print(f"WhatsApp Error: {e}")
        return False

def create_document_file(content: str, doc_type: str):
    buffer = io.BytesIO()
    doc_type = doc_type.lower()
    
    if doc_type == "docx":
        doc = Document()
        doc.add_heading("Dokumen Resmi - 37 AI Agents HQ", level=1)
        for line in content.split("\n"):
            doc.add_paragraph(line)
        doc.save(buffer)
        buffer.seek(0)
        return buffer, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "laporan_hq.docx"
    elif doc_type == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Data HQ"
        ws.append(["No", "Baris Konten / Analisis AI"])
        for idx, line in enumerate(content.split("\n"), start=1):
            if line.strip():
                ws.append([idx, line])
        wb.save(buffer)
        buffer.seek(0)
        return buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "data_hq.xlsx"
    elif doc_type == "pdf":
        p = canvas.Canvas(buffer, pagesize=letter)
        p.drawString(72, 750, "Dokumen Resmi - 37 AI Agents HQ")
        y = 720
        for line in content.split("\n"):
            if y < 50:
                p.showPage()
                y = 750
            p.drawString(72, y, line[:95])
            y -= 20
        p.save()
        buffer.seek(0)
        return buffer, "application/pdf", "laporan_hq.pdf"
    else:
        buffer.write(content.encode("utf-8"))
        buffer.seek(0)
        return buffer, "text/plain", "catatan_hq.txt"

@app.post("/api/chat")
async def chat(req: ChatRequest):
    try:
        division_style = DIVISION_PROMPTS.get(
            req.division, 
            "Bicara profesional dan solutif sesuai standar komunikasi kantor."
        )

        mastermind_instruction = f"""
        Kamu adalah {req.agent_name}, seorang {req.role} di divisi {req.division}. 
        Khusus sebagai Mastermind Ide Brilian Bos Septi:
        1. Lakukan pencarian internet secara aktif untuk tren iklan dan referensi visual terbaru jika diperlukan.
        2. Berikan ide prompt video iklan yang sangat sinematik, estetis, detail (camera angle, lighting, motion, skrip voiceover).
        3. Berikan jawaban yang terstruktur dengan tajam dan langsung siap dieksekusi.
        Gaya Komunikasi: {division_style}
        """

        ai_text = ""
        use_provider = req.provider.lower() if req.provider else "groq"

        # Prioritas Eksekusi via Groq (Mendukung Teks maupun Gambar via Qwen Vision)
        if use_provider == "groq" and groq_client:
            try:
                if req.image_base64:
                    header, encoded = req.image_base64.split(",", 1) if "," in req.image_base64 else ("", req.image_base64)
                    mime_type = "image/png" if "png" in header else "image/jpeg"
                    image_url = f"data:{mime_type};base64,{encoded}"
                    
                    messages = [
                        {"role": "system", "content": mastermind_instruction},
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": req.message},
                                {"type": "image_url", "image_url": {"url": image_url}}
                            ]
                        }
                    ]
                    groq_model = "qwen/qwen3.8-27b"
                else:
                    messages = [
                        {"role": "system", "content": mastermind_instruction},
                        {"role": "user", "content": req.message}
                    ]
                    groq_model = "openai/gpt-oss-120b"

                completion = groq_client.chat.completions.create(
                    model=groq_model,
                    messages=messages,
                    temperature=0.7,
                    max_tokens=1024
                )
                ai_text = completion.choices[0].message.content or ""
            except Exception as groq_err:
                print(f"⚠️ Groq Error (Fallback otomatis ke Gemini): {groq_err}")
                use_provider = "gemini"

        # Fallback otomatis ke Gemini jika Groq gagal atau provider diset ke gemini
        if (use_provider == "gemini" or not ai_text) and gemini_client:
            try:
                contents = [req.message]
                if req.image_base64:
                    header, encoded = req.image_base64.split(",", 1) if "," in req.image_base64 else ("", req.image_base64)
                    mime_type = "image/png" if "png" in header else "image/jpeg"
                    image_bytes = base64.b64decode(encoded)
                    contents.append(types.Part.from_bytes(data=image_bytes, mime_type=mime_type))

                response = gemini_client.models.generate_content(
                    model="gemini-3.5-flash-lite",
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=mastermind_instruction,
                        tools=[types.Tool(google_search=types.GoogleSearch())]
                    )
                )
                ai_text = response.text or ""
            except Exception as gemini_err:
                print(f"⚠️ Gemini Error: {gemini_err}")
                if "429" in str(gemini_err) or "RESOURCE_EXHAUSTED" in str(gemini_err):
                    ai_text = "⚠️ Kuota Gemini API saat ini sedang penuh (429 Rate Limit Exceeded). Mohon tunggu beberapa saat."
                else:
                    ai_text = f"⚠️ Terjadi kendala pada pemrosesan AI: {str(gemini_err)}"

        if not ai_text:
            ai_text = "Maaf Bos Septi, respon AI kosong. Silakan ulangi."

        log_agent_activity(req.agent_id, req.agent_name, req.division, ai_text)

        if req.auto_dispatch and not ai_text.startswith("⚠️"):
            send_to_telegram(ai_text)
            send_to_whatsapp(ai_text)

        voice = "id-ID-GadisNeural" if req.gender.lower() == "female" else "id-ID-ArdiNeural"
        communicate = edge_tts.Communicate(ai_text, voice)
        audio_bytes = b""
        async for chunk in communicate.stream():
            if chunk["type"] == "audio":
                audio_bytes += chunk["data"]

        audio_base64 = base64.b64encode(audio_bytes).decode('utf-8')

        return {
            "response": ai_text,
            "audio": audio_base64
        }

    except Exception as e:
        print(f"Error Chat Endpoint: {e}")
        raise HTTPException(status_code=500, detail="Gagal memproses pesan AI Mastermind.")

@app.post("/api/generate-document")
async def generate_document(req: DocRequest):
    try:
        buffer, media_type, filename = create_document_file(req.content, req.doc_type)
        return StreamingResponse(
            buffer,
            media_type=media_type,
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        print(f"Error Doc Gen: {e}")
        raise HTTPException(status_code=500, detail="Gagal men-generate file dokumen.")

@app.post("/api/video-pipeline/trigger")
async def trigger_pipeline_manual(background_tasks: BackgroundTasks):
    """Trigger manual pengerjaan 3 video short dari Dashboard Canvas 3D"""
    log_agent_activity(
        "1", 
        "Aria", 
        "Mastermind", 
        "Memulai eksekusi otomatisasi pembuatan 3 video short sinematik atas perintah Bos Septi."
    )
    background_tasks.add_task(run_short_production_pipeline)
    return {"status": "success", "message": "Pipeline otomatisasi 3 video short telah dimulai oleh tim agen."}